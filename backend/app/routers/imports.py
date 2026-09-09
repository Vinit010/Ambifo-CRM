import re
from datetime import date, datetime, timezone
from io import BytesIO

import requests
from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile, status
from sqlalchemy import func, or_
from sqlalchemy.orm import Session

from ..database import get_db
from ..deps import get_current_user
from ..models.crm import Customer, Lead, OpportunityHistory
from ..models.user import User
from ..schemas.engine import (
    CsvAnalyzeIn,
    CsvImportIn,
    CsvPreviewOut,
    ExcelImportIn,
    ImportResultOut,
)
from ..services.engine_client import engine_client

router = APIRouter(prefix="/api/import", tags=["imports"])

# Source header (normalized) -> Customer/Lead attribute
CUSTOMER_ALIASES = {
    "customer_name": "customer_name",
    "name": "customer_name",
    "customer": "customer_name",
    "company": "account_name",
    "account": "account_name",
    "account_name": "account_name",
    "email": "email",
    "emails": "email",
    "phone": "phone",
    "contact": "phone",
    "phone_number": "phone",
    "city": "city",
    "cloud": "cloud",
    "cloud_operator": "cloud",
    "segment": "segment",
    "deal_status": "deal_status",
    "status": "deal_status",
    "aws_id": "aws_id",
    "aws": "aws_id",
    "aws_account_id": "aws_id",
    "opportunity_id": "opportunity_id",
    "opp_id": "opportunity_id",
    "comment": "comment",
    "notes": "comment",
    "main_page": "main_page_address",
    "website": "main_page_address",
    "billing": "billing",
}

LEAD_ALIASES = {
    "lead_name": "lead_name",
    "name": "lead_name",
    "lead": "lead_name",
    "email": "email",
    "phone": "phone",
    "company": "company",
    "city": "city",
    "source": "source",
    "tags": "tags",
    "notes": "notes",
    "status": "lead_status",
}


def csv_rows_to_records(
    headers: list[str], rows: list[list[str]], aliases: dict[str, str], column_map: dict[str, str]
) -> list[dict]:
    """Convert CSV output into a list of entity dicts for known fields."""
    records: list[dict] = []
    for row in rows:
        attrs: dict[str, str] = {}
        for idx, header in enumerate(headers):
            norm = header.strip().lower().replace(" ", "_").replace("-", "_")
            target = column_map.get(header, aliases.get(norm))
            if not target:
                continue
            if idx < len(row):
                val = row[idx].strip()
                if val and val.lower() not in ("n/a", "null", "na"):
                    attrs[target] = val
        records.append(attrs)
    return records


def now() -> datetime:
    return datetime.now(timezone.utc)


CUSTOMER_UPDATABLE = {
    "account_name",
    "customer_name",
    "designation",
    "email",
    "phone",
    "cloud",
    "main_page_address",
    "billing",
    "city",
    "aws_id",
    "aws_calculator_link",
    "opportunity_id",
    "segment",
    "deal_status",
    "comment",
}

_DOWNLOAD_RE = re.compile(r"/_layouts/15/(?:Doc|doc)\d*\.aspx")


def _to_download_url(url: str) -> str:
    """Convert a SharePoint web link into a direct download URL."""
    cleaned = re.sub(r"&amp;", "&", url.strip())
    cleaned = _DOWNLOAD_RE.sub("/_layouts/15/download.aspx", cleaned)
    cleaned = cleaned.split("&action=")[0]
    if "?" not in cleaned:
        cleaned += "?download=1"
    return cleaned


def _download_workbook(url: str) -> bytes:
    try:
        resp = requests.get(
            url,
            timeout=60,
            allow_redirects=True,
            headers={
                "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
            },
        )
        resp.raise_for_status()
    except requests.RequestException as exc:
        raise HTTPException(
            status_code=502,
            detail=f"Could not download workbook from link ({exc}). Make sure the file is shared with 'Anyone with the link'.",
        ) from exc
    return resp.content


def _cell_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value).strip()


def _read_workbook(content: bytes, sheet_name: str | None) -> tuple[list[str], list[list[str]]]:
    import openpyxl

    wb = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb[wb.sheetnames[0]]
    header_row = None
    rows: list[list[str]] = []
    for row in ws.iter_rows(values_only=True):
        cells = [_cell_str(v) for v in row]
        if not any(cells):
            continue
        if header_row is None:
            header_row = cells
        else:
            rows.append(cells)
    wb.close()
    return header_row or [], rows


def _synthetic_email(name: str, row_idx: int) -> str:
    slug = re.sub(r"[^a-z0-9]+", ".", name.lower()).strip(".")[:40]
    return f"{slug or 'importrow'}{row_idx}@import.ambifo.invalid"


def _content_to_table(
    content: bytes, filename: str, sheet_name: str | None
) -> tuple[list[str], list[list[str]]]:
    """Parse an uploaded file into (headers, rows). Supports Excel and CSV."""
    lowered = (filename or "").lower()
    if lowered.endswith((".csv", ".txt")):
        text = content.decode("utf-8-sig", errors="replace")
        parsed = engine_client.parse_csv(text, True)
        return parsed.get("headers", []), parsed.get("rows", [])
    return _read_workbook(content, sheet_name)


def _upsert_customers(
    db: Session,
    headers: list[str],
    rows: list[list[str]],
    *,
    update_existing: bool,
    column_map: dict[str, str],
    user: User,
    source: str,
) -> ImportResultOut:
    records = csv_rows_to_records(headers, rows, CUSTOMER_ALIASES, column_map)
    created_ids: list[int] = []
    errors: list[str] = []
    created = updated = skipped = 0

    for idx, rec in enumerate(records, start=2):
        name = (rec.get("customer_name") or "").strip()
        email = (rec.get("email") or "").strip().lower()
        if not name and not email:
            errors.append(f"row {idx}: missing name and email")
            continue

        exists = None
        if email:
            exists = db.query(Customer).filter(func.lower(Customer.email) == email).first()
        if not exists and name:
            exists = db.query(Customer).filter(func.lower(Customer.customer_name) == name.lower()).first()

        fields = {k: v for k, v in rec.items() if k in CUSTOMER_UPDATABLE}

        if not exists:
            if not email:
                email = _synthetic_email(name, idx)
            attrs = {k: v for k, v in fields.items() if k != "customer_name"}
            customer = Customer(customer_name=name or email.split("@")[0], email=email, **attrs)
            db.add(customer)
            db.flush()
            hist = OpportunityHistory()
            hist.customer_id = customer.id
            hist.action = "created"
            hist.changes_summary = f"Imported from {source} (row {idx})"
            hist.changed_by = user.username
            hist.created_at = now()
            db.add(hist)
            created += 1
            created_ids.append(customer.id)
            continue

        if not update_existing:
            skipped += 1
            continue

        for key, val in fields.items():
            setattr(exists, key, val)
        if email:
            exists.email = email
        if name and name.lower() != (exists.customer_name or "").lower():
            exists.customer_name = name

        hist = OpportunityHistory()
        hist.customer_id = exists.id
        hist.action = "updated"
        hist.changes_summary = f"Updated from {source} (row {idx}): {', '.join(sorted(fields))}"
        hist.changed_by = user.username
        hist.created_at = now()
        db.add(hist)
        updated += 1

    db.commit()
    return ImportResultOut(
        entity="customers",
        created=created,
        created_ids=created_ids,
        skipped=skipped,
        updated=updated,
        errors=errors,
    )


@router.post("/csv-analyze", response_model=CsvPreviewOut)
def analyze_csv(
    payload: CsvAnalyzeIn,
    _user: User = Depends(get_current_user),
):
    """Send raw CSV to the Rust engine and return a preview (headers + rows)."""
    try:
        result = engine_client.parse_csv(payload.content, payload.contains_header)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=502, detail=f"Engine CSV parse failed: {exc}") from exc
    return CsvPreviewOut(
        headers=result.get("headers", []),
        rows=result.get("rows", []),
        row_count=result.get("row_count", 0),
    )


@router.post("/customers", response_model=ImportResultOut, status_code=status.HTTP_201_CREATED)
def import_customers(
    payload: CsvImportIn,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    if payload.entity != "customers":
        raise HTTPException(status_code=422, detail="Use /api/import/leads for leads")
    try:
        parsed = engine_client.parse_csv(payload.content, payload.contains_header)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=502, detail=f"Engine CSV parse failed: {exc}") from exc

    headers = parsed.get("headers", [])
    if not headers:
        raise HTTPException(status_code=422, detail="CSV has no headers; provide them or set contains_header")

    records = csv_rows_to_records(headers, parsed.get("rows", []), CUSTOMER_ALIASES, payload.column_map)
    created_ids: list[int] = []
    errors: list[str] = []
    created = skipped = 0
    for idx, rec in enumerate(records, start=1):
        if not rec.get("customer_name"):
            errors.append(f"row {idx}: missing customer name")
            continue
        email = (rec.get("email") or "").lower().strip()
        exists = None
        if email:
            exists = (
                db.query(Customer)
                .filter(or_(Customer.email == email, Customer.customer_name == rec["customer_name"]))
                .first()
            )
        else:
            exists = (
                db.query(Customer)
                .filter(Customer.customer_name == rec["customer_name"])
                .first()
            )
        if exists:
            skipped += 1
            continue
        customer = Customer(customer_name=rec["customer_name"], **{k: v for k, v in rec.items() if k != "customer_name"})
        db.add(customer)
        db.flush()
        hist = OpportunityHistory()
        hist.customer_id = customer.id
        hist.action = "created"
        hist.changes_summary = f"Imported from CSV (row {idx})"
        hist.changed_by = current_user.username
        hist.created_at = now()
        db.add(hist)
        created += 1
        created_ids.append(customer.id)

    db.commit()
    return ImportResultOut(entity="customers", created=created, created_ids=created_ids, skipped=skipped, errors=errors)


@router.post("/excel-customers", response_model=ImportResultOut, status_code=status.HTTP_201_CREATED)
def import_excel_customers(
    payload: ExcelImportIn,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Download an Excel workbook (e.g. a SharePoint/OneDrive link) and upsert customers."""
    try:
        content = _download_workbook(_to_download_url(payload.url))
    except HTTPException:
        raise
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=502, detail=f"Workbook download failed: {exc}") from exc

    try:
        headers, rows = _read_workbook(content, payload.sheet_name)
    except HTTPException:
        raise
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=422, detail=f"Could not read Excel workbook: {exc}") from exc

    if not headers:
        raise HTTPException(status_code=422, detail="Workbook has no header row")

    return _upsert_customers(
        db,
        headers,
        rows,
        update_existing=payload.update_existing,
        column_map=payload.column_map,
        user=current_user,
        source="Excel link",
    )


@router.post("/customers/excel", response_model=ImportResultOut, status_code=status.HTTP_201_CREATED)
async def upload_customers_excel(
    update_existing: bool = Query(default=True),
    sheet_name: str | None = Query(default=None),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Upsert customers from an uploaded Excel/CSV file."""
    content = await file.read()
    if not content:
        raise HTTPException(status_code=422, detail="Empty file")

    filename = file.filename or "upload"
    try:
        headers, rows = _content_to_table(content, filename, sheet_name)
    except HTTPException:
        raise
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=422, detail=f"Could not read file: {exc}") from exc

    if not headers:
        raise HTTPException(status_code=422, detail="File has no header row")

    return _upsert_customers(
        db,
        headers,
        rows,
        update_existing=update_existing,
        column_map={},
        user=current_user,
        source=f"Uploaded file '{filename}'",
    )


@router.post("/leads", response_model=ImportResultOut, status_code=status.HTTP_201_CREATED)
def import_leads(
    payload: CsvImportIn,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    if payload.entity != "leads":
        raise HTTPException(status_code=422, detail="Use /api/import/customers for customers")
    try:
        parsed = engine_client.parse_csv(payload.content, payload.contains_header)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=502, detail=f"Engine CSV parse failed: {exc}") from exc

    headers = parsed.get("headers", [])
    if not headers:
        raise HTTPException(status_code=422, detail="CSV has no headers; provide them or set contains_header")

    records = csv_rows_to_records(headers, parsed.get("rows", []), LEAD_ALIASES, payload.column_map)
    created_ids: list[int] = []
    errors: list[str] = []
    created = skipped = 0
    for idx, rec in enumerate(records, start=1):
        if not rec.get("lead_name") and not rec.get("email"):
            errors.append(f"row {idx}: missing name and email")
            continue
        email = (rec.get("email") or "").lower().strip()
        exists = None
        if email:
            exists = db.query(Lead).filter(Lead.email == email).first()
        if exists:
            skipped += 1
            continue
        lead = Lead(**{k: (v.lower() if k == "lead_status" else v) for k, v in rec.items()})
        db.add(lead)
        db.flush()
        created += 1
        created_ids.append(lead.id)

    db.commit()
    return ImportResultOut(entity="leads", created=created, created_ids=created_ids, skipped=skipped, errors=errors)